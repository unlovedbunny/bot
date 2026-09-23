from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def criar_template_basico():
    doc = Document()
    
    # Título
    titulo = doc.add_heading('Meu Nome Completo', 0)
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Info de Contato
    contato = doc.add_paragraph('email@exemplo.com | (11) 99999-9999 | LinkedIn: /in/meuperfil')
    contato.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Resumo (Aqui entra a tag da IA)
    doc.add_heading('Resumo Profissional', level=1)
    # A TAG ABAIXO É O QUE O NOSSO SCRIPT VAI BUSCAR E SUBSTITUIR
    p_resumo = doc.add_paragraph('{{RESUMO_IA}}')
    
    # Experiência
    doc.add_heading('Experiência Profissional', level=1)
    
    exp1 = doc.add_paragraph()
    exp1.add_run('Engenheiro de Software Sênior - Empresa X\n').bold = True
    exp1.add_run('Jan/2020 - O Presente\n')
    exp1.add_run('- Desenvolvimento de pipelines de dados em Python.\n')
    exp1.add_run('- Automação de processos internos com Playwright.\n')
    
    # Educação
    doc.add_heading('Formação Acadêmica', level=1)
    edu = doc.add_paragraph()
    edu.add_run('Bacharelado em Ciência da Computação - Universidade Y\n').bold = True
    edu.add_run('2015 - 2019')
    
    # Salva o documento
    doc.save('curriculo_template.docx')
    print("Template 'curriculo_template.docx' criado com sucesso!")

if __name__ == "__main__":
    criar_template_basico()
